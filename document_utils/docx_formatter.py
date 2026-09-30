from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from backend.config import COMPANY_NAME, COMPANY_TAGLINE
from document_utils.sanitize import sanitize_text


def format_docx(text: str, doc_type: str) -> bytes:
    """
    Convert generated legal document text into a DOCX file.
    """

    document = Document()

    # ---------------------------------------------------------
    # Page margins
    # ---------------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # ---------------------------------------------------------
    # Default font
    # ---------------------------------------------------------

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    # ---------------------------------------------------------
    # Header / Brand
    # ---------------------------------------------------------

    brand = document.add_paragraph()

    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER

    brand_run = brand.add_run(
        f"{COMPANY_NAME}"
    )

    brand_run.bold = True
    brand_run.font.name = "Times New Roman"
    brand_run.font.size = Pt(18)

    # ---------------------------------------------------------
    # Document title
    # ---------------------------------------------------------

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(
        sanitize_text(doc_type).upper()
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    # ---------------------------------------------------------
    # Tagline
    # ---------------------------------------------------------

    subtitle = document.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle_run = subtitle.add_run(
        COMPANY_TAGLINE
    )

    subtitle_run.italic = True
    subtitle_run.font.name = "Times New Roman"
    subtitle_run.font.size = Pt(9)

    document.add_paragraph()

    # ---------------------------------------------------------
    # Document body
    # ---------------------------------------------------------

    body = sanitize_text(text)

    for raw_line in body.splitlines():

        line = raw_line.strip()

        # Empty line
        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15

        # Detect headings
        is_heading = (
            len(line) < 120
            and (
                line.isupper()
                or line.endswith(":")
                or (
                    len(line) >= 2
                    and line[0].isdigit()
                    and line[1] == "."
                )
            )
        )

        if is_heading:

            run = paragraph.add_run(line)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

        else:

            run = paragraph.add_run(line)

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    # ---------------------------------------------------------
    # Footer
    # ---------------------------------------------------------

    footer = section.footer.paragraphs[0]

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer.add_run(
        f"{COMPANY_NAME} | AI Generated Draft | Review Before Use"
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    # ---------------------------------------------------------
    # Save DOCX to memory
    # ---------------------------------------------------------

    output = BytesIO()

    document.save(output)

    return output.getvalue()